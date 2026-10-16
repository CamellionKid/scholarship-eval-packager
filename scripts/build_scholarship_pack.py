#!/usr/bin/env python3
"""
Build a complete university scholarship self-evaluation package:
1. Standardized attachment directory + original blank form backup
2. Formula-driven 3-sheet Excel workbook (.xlsx) with clickable attachment links
3. Filled Word form (.docx & .doc)
4. Optional PDF/PNG visual rendering via LibreOffice + pdftoppm
"""

import argparse
import json
import os
import shutil
import subprocess
from collections import OrderedDict

import docx
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def copy_attachments(base_dir: str, copies: list[dict], word_template: str | None) -> None:
    if word_template and os.path.exists(word_template):
        backup_dir = os.path.join(base_dir, "附件_支撑材料", "00_原始空白模板备份")
        os.makedirs(backup_dir, exist_ok=True)
        stem, ext = os.path.splitext(os.path.basename(word_template))
        backup_path = os.path.join(backup_dir, f"{stem}_空白原件{ext}")
        if not os.path.exists(backup_path):
            shutil.copy2(word_template, backup_path)

    for entry in copies:
        src = entry.get("src")
        dst_rel = entry.get("dst_rel")
        if not src or not dst_rel:
            continue
        dst = os.path.join(base_dir, dst_rel)
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        if os.path.exists(src):
            shutil.copy2(src, dst)


def build_excel(spec: dict, excel_path: str) -> None:
    wb = openpyxl.Workbook()

    font_title = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
    font_subtitle = Font(name="微软雅黑", size=10, italic=True, color="595959")
    font_header = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    font_section = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")
    font_body = Font(name="微软雅黑", size=10, bold=False, color="000000")
    font_bold = Font(name="微软雅黑", size=10, bold=True, color="000000")
    font_link = Font(name="微软雅黑", size=9, underline="single", color="0563C1")
    font_note = Font(name="微软雅黑", size=9, color="595959")

    fill_header = PatternFill("solid", fgColor="1F4E79")
    fill_subhdr = PatternFill("solid", fgColor="D9E1F2")
    fill_input = PatternFill("solid", fgColor="FFF2CC")
    fill_formula = PatternFill("solid", fgColor="E2EFDA")
    fill_total = PatternFill("solid", fgColor="FCE4D6")

    thin_side = Side(border_style="thin", color="BFBFBF")
    med_side = Side(border_style="medium", color="1F4E79")
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_total = Border(left=thin_side, right=thin_side, top=med_side, bottom=med_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    app = spec["applicant"]
    courses = spec.get("courses", [])
    course_end_row = 3 + len(courses)
    round_avg_row = course_end_row + 3

    # ------------------------------------------------------------------
    # Sheet 1: 1_分项自动计分明细表
    # ------------------------------------------------------------------
    ws_calc = wb.active
    ws_calc.title = "1_分项自动计分明细表"
    ws_calc.views.sheetView[0].showGridLines = True

    ws_calc.merge_cells("A1:J1")
    ws_calc["A1"] = spec.get("title", "奖学金自评自动核算表")
    ws_calc["A1"].font = font_title
    ws_calc["A1"].alignment = align_center
    ws_calc.row_dimensions[1].height = 32

    ws_calc.merge_cells("A2:J2")
    ws_calc["A2"] = (
        f"申请人：{app['name']}  |  学号：{app['id']}  |  年级专业：{app.get('major', '')}  |  "
        f"评审依据：{app.get('rule_doc', '')}  |  说明：黄色单元格为原始分输入项，绿色/橙色为Excel公式自动计算结果"
    )
    ws_calc["A2"].font = font_subtitle
    ws_calc["A2"].alignment = align_center
    ws_calc.row_dimensions[2].height = 22

    headers_calc = [
        "序号",
        "一级指标（权重）",
        "对应自评表栏目",
        "成果/项目明细与计算说明",
        "评审细则条款依据",
        "原始分值\n(可改输入)",
        "指标权重",
        "折算自评分\n(=原始分×权重)",
        "支撑材料文件（点击可跳转）",
        "备注说明",
    ]
    ws_calc.row_dimensions[3].height = 28
    for col_idx, h in enumerate(headers_calc, 1):
        c = ws_calc.cell(row=3, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_all

    items = spec.get("items", [])
    col_key_rows = OrderedDict()

    for idx, item in enumerate(items, start=4):
        ws_calc.row_dimensions[idx].height = item.get("height", 46)
        raw_val = item["raw_score"]
        if raw_val == "__COURSE_AVG__":
            raw_val = f"='3_课程成绩明细表'!E{round_avg_row}"

        c1 = ws_calc.cell(row=idx, column=1, value=item["no"])
        c2 = ws_calc.cell(row=idx, column=2, value=item["cat"])
        c3 = ws_calc.cell(row=idx, column=3, value=item["col_name"])
        c4 = ws_calc.cell(row=idx, column=4, value=item["detail"])
        c5 = ws_calc.cell(row=idx, column=5, value=item["rule"])
        c6 = ws_calc.cell(row=idx, column=6, value=raw_val)
        c7 = ws_calc.cell(row=idx, column=7, value=item["weight"])
        c8 = ws_calc.cell(row=idx, column=8, value=f"=F{idx}*G{idx}")
        c9 = ws_calc.cell(row=idx, column=9, value=item.get("link_text", "无"))
        c10 = ws_calc.cell(row=idx, column=10, value=item.get("note", ""))

        for c in (c1, c2, c3, c4, c5, c6, c7, c8, c9, c10):
            c.border = border_all
            c.font = font_body

        c1.alignment = align_center
        c2.alignment = align_center
        c2.font = font_section
        c3.alignment = align_center
        c4.alignment = align_left
        c5.alignment = align_left
        c6.alignment = align_center
        c6.fill = fill_input
        c6.font = font_bold
        c6.number_format = "0.000" if item.get("raw_score") == "__COURSE_AVG__" else "0.##"
        c7.alignment = align_center
        c7.number_format = "0%"
        c8.alignment = align_center
        c8.fill = fill_formula
        c8.font = font_bold
        c8.number_format = "0.0000"

        if item.get("link_path"):
            c9.hyperlink = item["link_path"]
            c9.font = font_link
        c9.alignment = align_left
        c10.alignment = align_left
        c10.font = font_note

        ck = item.get("col_key", item["col_name"])
        col_key_rows.setdefault(ck, {"label": item["col_name"].split("\n")[0], "rows": []})["rows"].append(idx)

    last_item_row = 3 + len(items)
    start_sum_row = last_item_row + 1
    col_key_summary_row = {}

    for offset, (ck, info) in enumerate(col_key_rows.items()):
        r = start_sum_row + offset
        col_key_summary_row[ck] = r
        r_list = info["rows"]
        if len(r_list) == 1:
            raw_f = f"=F{r_list[0]}"
            w_f = f"=H{r_list[0]}"
        else:
            raw_f = f"=SUM(F{r_list[0]}:F{r_list[-1]})"
            w_f = f"=SUM(H{r_list[0]}:H{r_list[-1]})"

        ws_calc.row_dimensions[r].height = 24
        ws_calc.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws_calc.cell(row=r, column=1, value=f"模块小计：{info['label']}").alignment = align_right
        ws_calc.cell(row=r, column=6, value=raw_f).alignment = align_center
        ws_calc.cell(row=r, column=7, value="—").alignment = align_center
        ws_calc.cell(row=r, column=8, value=w_f).alignment = align_center
        ws_calc.merge_cells(start_row=r, start_column=9, end_row=r, end_column=10)
        ws_calc.cell(row=r, column=9, value=f"对应自评表栏目：{info['label']}").alignment = align_left

        for c_i in range(1, 11):
            cell = ws_calc.cell(row=r, column=c_i)
            cell.border = border_all
            cell.fill = fill_subhdr
            cell.font = font_bold if c_i in (1, 6, 8) else font_note
        ws_calc.cell(row=r, column=6).number_format = "0.###"
        ws_calc.cell(row=r, column=8).number_format = "0.0000"

    tot_r = start_sum_row + len(col_key_rows)
    ws_calc.row_dimensions[tot_r].height = 30
    ws_calc.merge_cells(start_row=tot_r, start_column=1, end_row=tot_r, end_column=5)
    ws_calc.cell(row=tot_r, column=1, value="自评总分合计（自动求和）").alignment = align_right
    ws_calc.cell(row=tot_r, column=6, value="—").alignment = align_center
    ws_calc.cell(row=tot_r, column=7, value="100%").alignment = align_center
    ws_calc.cell(row=tot_r, column=8, value=f"=SUM(H4:H{last_item_row})").alignment = align_center
    ws_calc.merge_cells(start_row=tot_r, start_column=9, end_row=tot_r, end_column=10)
    ws_calc.cell(row=tot_r, column=9, value=f"计算公式：=SUM(H4:H{last_item_row})").alignment = align_left

    for c_i in range(1, 11):
        cell = ws_calc.cell(row=tot_r, column=c_i)
        cell.border = border_total
        cell.fill = fill_total
        cell.font = Font(name="微软雅黑", size=11, bold=True, color="C00000" if c_i == 8 else "000000")
    ws_calc.cell(row=tot_r, column=8).number_format = "0.0000"

    for i, w in enumerate([6, 16, 20, 44, 34, 13, 10, 15, 36, 32], 1):
        ws_calc.column_dimensions[get_column_letter(i)].width = w

    # ------------------------------------------------------------------
    # Sheet 2: 2_学院标准自评分表(可打印)
    # ------------------------------------------------------------------
    ws_form = wb.create_sheet(title="2_学院标准自评分表(可打印)")
    ws_form.views.sheetView[0].showGridLines = True

    word_cols = spec.get("word_columns", [])
    num_cols = len(word_cols)
    last_col_letter = get_column_letter(num_cols)

    ws_form.merge_cells(f"A1:{last_col_letter}1")
    ws_form["A1"] = spec.get("form_title", "硕士研究生学业奖学金自评分表")
    ws_form["A1"].font = font_title
    ws_form["A1"].alignment = align_center
    ws_form.row_dimensions[1].height = 34

    ws_form.row_dimensions[2].height = 30
    for col_idx, col_def in enumerate(word_cols, 1):
        c = ws_form.cell(row=2, column=col_idx, value=col_def["header"])
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_all

    ws_form.row_dimensions[3].height = 175
    ws_form.row_dimensions[4].height = 38
    row_texts = spec.get("word_row_texts", {})

    score_col_indices = []
    for col_idx, col_def in enumerate(word_cols, 1):
        key = col_def["key"]
        val3 = f"='1_分项自动计分明细表'!H{tot_r}" if key == "total" else row_texts.get(key, "")
        c3 = ws_form.cell(row=3, column=col_idx, value=val3)
        c3.border = border_all
        c3.font = font_bold if key in ("id", "name", "total") else font_body
        c3.alignment = align_center if col_def.get("align") == "center" else align_left
        if key == "total":
            c3.fill = fill_total
            c3.number_format = "0.0000"

        if key == "id":
            val4 = app["id"]
        elif key == "name":
            val4 = f"{app['name']}（纯数值公式核验行）"
        elif key in col_key_summary_row:
            val4 = f"='1_分项自动计分明细表'!H{col_key_summary_row[key]}"
            score_col_indices.append(col_idx)
        elif key == "total" and score_col_indices:
            first_c = get_column_letter(score_col_indices[0])
            last_c = get_column_letter(score_col_indices[-1])
            val4 = f"=SUM({first_c}4:{last_c}4)"
        else:
            val4 = ""

        c4 = ws_form.cell(row=4, column=col_idx, value=val4)
        c4.border = border_total
        c4.fill = fill_total if key == "total" else fill_formula
        c4.font = font_bold
        c4.alignment = align_center
        if key in col_key_summary_row or key == "total":
            c4.number_format = "0.0000"

        ws_form.column_dimensions[get_column_letter(col_idx)].width = max(14, int(col_def.get("width_cm", 2.0) * 6.5))

    # ------------------------------------------------------------------
    # Sheet 3: 3_课程成绩明细表
    # ------------------------------------------------------------------
    ws_grade = wb.create_sheet(title="3_课程成绩明细表")
    ws_grade.views.sheetView[0].showGridLines = True

    ws_grade.merge_cells("A1:F1")
    ws_grade["A1"] = "研究生成绩单明细与学业成绩（30%）自动核算表"
    ws_grade["A1"].font = font_title
    ws_grade["A1"].alignment = align_center
    ws_grade.row_dimensions[1].height = 30

    ws_grade.merge_cells("A2:F2")
    ws_grade["A2"] = (
        f"姓名：{app['name']}  |  学号：{app['id']}  |  所在院系：{app.get('department', '')}  |  "
        f"专业：{app.get('major', '')}  |  导师：{app.get('advisor', '')}"
    )
    ws_grade["A2"].font = font_subtitle
    ws_grade["A2"].alignment = align_center
    ws_grade.row_dimensions[2].height = 20

    headers_grade = ["序号", "课程编号", "课程名称", "学分", "成绩（可改输入）", "备注/数据来源"]
    ws_grade.row_dimensions[3].height = 26
    for col_idx, h in enumerate(headers_grade, 1):
        c = ws_grade.cell(row=3, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_all

    for idx, course in enumerate(courses, start=4):
        ws_grade.row_dimensions[idx].height = 22
        vals = [
            course["no"],
            course["code"],
            course["name"],
            course["credit"],
            course["score"],
            course.get("source", "盖章成绩单 & 学院平均成绩公示表"),
        ]
        for c_i, v in enumerate(vals, 1):
            c = ws_grade.cell(row=idx, column=c_i, value=v)
            c.border = border_all
            c.font = font_bold if c_i == 5 else font_body
            c.alignment = align_left if c_i in (3, 6) else align_center
            if c_i == 4:
                c.number_format = "0.0"
            elif c_i == 5:
                c.fill = fill_input
                c.number_format = "0.00"

    r_sum = course_end_row + 1
    r_avg = course_end_row + 2
    r_ravg = course_end_row + 3
    r_w1 = course_end_row + 4
    r_w2 = course_end_row + 5

    grade_summaries = [
        (r_sum, f"课程总学分 / {len(courses)}门课程总分", f"=SUM(D4:D{course_end_row})", f"=SUM(E4:E{course_end_row})", "0.00", f"公式：=SUM(E4:E{course_end_row})"),
        (r_avg, "算术平均分（未舍入原始值）", "—", f"=AVERAGE(E4:E{course_end_row})", "0.000000", f"公式：=AVERAGE(E4:E{course_end_row})"),
        (r_ravg, "算术平均分（除不尽保留三位小数）", "—", f"=ROUND(AVERAGE(E4:E{course_end_row}), 3)", "0.000", f"公式：=ROUND(AVERAGE(E4:E{course_end_row}), 3)"),
        (r_w1, "学业成绩折算分（按保留三位小数平均分 × 30%）", "30%", f"=E{r_ravg}*0.3", "0.0000", f"公式：=E{r_ravg}*0.3"),
        (r_w2, "学业成绩折算分（按未舍入平均分 × 30% 对照）", "30%", f"=E{r_avg}*0.3", "0.0000", f"公式：=E{r_avg}*0.3"),
    ]

    for row_num, label, d_val, e_formula, num_fmt, remark in grade_summaries:
        ws_grade.row_dimensions[row_num].height = 24
        ws_grade.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
        ws_grade.cell(row=row_num, column=1, value=label).alignment = align_right
        ws_grade.cell(row=row_num, column=4, value=d_val).alignment = align_center
        ws_grade.cell(row=row_num, column=5, value=e_formula).alignment = align_center
        ws_grade.cell(row=row_num, column=6, value=remark).alignment = align_left
        for c_i in range(1, 7):
            c = ws_grade.cell(row=row_num, column=c_i)
            c.border = border_total if row_num == r_w1 else border_all
            c.fill = fill_total if row_num == r_w1 else fill_subhdr
            c.font = font_bold if c_i in (1, 5) else font_note
        ws_grade.cell(row=row_num, column=5).number_format = num_fmt

    for i, w in enumerate([8, 16, 38, 10, 18, 42], 1):
        ws_grade.column_dimensions[get_column_letter(i)].width = w

    for ws in wb.worksheets:
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1

    wb.save(excel_path)


def fill_word_form(spec: dict, base_dir: str) -> str | None:
    word_template = spec.get("word_template")
    if not word_template or not os.path.exists(word_template):
        return None

    tmp_dir = "/tmp/scholarship_pack_lo"
    os.makedirs(tmp_dir, exist_ok=True)
    stem, ext = os.path.splitext(os.path.basename(word_template))

    if ext.lower() == ".doc":
        subprocess.run(
            [
                "soffice",
                "-env:UserInstallation=file:///tmp/lo_profile_pack",
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                tmp_dir,
                word_template,
            ],
            check=True,
        )
        working_docx = os.path.join(tmp_dir, f"{stem}.docx")
    else:
        working_docx = word_template

    doc = docx.Document(working_docx)
    table = doc.tables[0]
    row = table.rows[1]
    word_cols = spec.get("word_columns", [])
    row_texts = spec.get("word_row_texts", {})

    for r in table.rows:
        for idx, col_def in enumerate(word_cols):
            if idx < len(r.cells):
                r.cells[idx].width = Cm(col_def.get("width_cm", 2.0))

    for idx, col_def in enumerate(word_cols):
        if idx >= len(row.cells):
            break
        key = col_def["key"]
        text = row_texts.get(key, "")
        cell = row.cells[idx]
        cell.text = ""
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        lines = text.split("\n") if text else [""]
        for l_i, line in enumerate(lines):
            p = cell.paragraphs[0] if l_i == 0 else cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if col_def.get("align") == "center"
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = p.add_run(line)
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(8.5 if key == "id" else 9.5)
            if key in ("id", "name", "total") or line.startswith("总计："):
                run.font.bold = True

    out_docx = os.path.join(base_dir, f"{stem}.docx")
    doc.save(out_docx)

    if ext.lower() == ".doc":
        subprocess.run(
            [
                "soffice",
                "-env:UserInstallation=file:///tmp/lo_profile_pack",
                "--headless",
                "--convert-to",
                "doc",
                "--outdir",
                base_dir,
                out_docx,
            ],
            check=True,
        )
    return out_docx


def render_previews(files: list[str], render_dir: str) -> None:
    os.makedirs(render_dir, exist_ok=True)
    for fpath in files:
        if not fpath or not os.path.exists(fpath):
            continue
        subprocess.run(
            [
                "soffice",
                "-env:UserInstallation=file:///tmp/lo_profile_pack",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                render_dir,
                fpath,
            ],
            check=True,
        )
        stem = os.path.splitext(os.path.basename(fpath))[0]
        pdf_path = os.path.join(render_dir, f"{stem}.pdf")
        if os.path.exists(pdf_path):
            subprocess.run(
                ["pdftoppm", "-png", "-r", "150", pdf_path, os.path.join(render_dir, stem)],
                check=True,
            )


def main() -> None:
    parser = argparse.ArgumentParser(description="Build scholarship self-evaluation Excel, Word form, and attachment pack.")
    parser.add_argument("--spec", required=True, help="Path to JSON specification file")
    parser.add_argument("--render-dir", default=None, help="Optional directory to render PDF/PNG previews")
    args = parser.parse_args()

    with open(args.spec, "r", encoding="utf-8") as f:
        spec = json.load(f)

    base_dir = spec["base_dir"]
    os.makedirs(base_dir, exist_ok=True)

    copy_attachments(base_dir, spec.get("copies", []), spec.get("word_template"))

    excel_path = os.path.join(base_dir, spec["excel_filename"])
    build_excel(spec, excel_path)
    out_docx = fill_word_form(spec, base_dir)

    if args.render_dir:
        render_previews([excel_path, out_docx], args.render_dir)

    print(f"Package built in: {base_dir}")


if __name__ == "__main__":
    main()
